"""Ledger explorer read model and Excel export. No production writes."""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

from arbicore.opportunity_ledger.export_xlsx import SHEETS, build_workbook
from arbicore.opportunity_ledger.read_model import (
    apply_filters,
    detail,
    paginate,
    slim_row,
    sort_rows,
    summarize,
    timeline_rows,
)

RUN = "shadowcert-5605e7b9-004b-49c7-ab8e-f618caba4065"
SECRET = re.compile(r"(://|mongodb|PRIVATE KEY|eyJ[A-Za-z0-9_-]{8,}\.)", re.I)


def _row(
    cid,
    family,
    net,
    *,
    chain="base",
    provider="balancer_v2",
    bundle=True,
    gate="FAIL",
    slippage=0.0,
    flash_fee=0.0,
    gross_profit=None,
    verified_at=1791178200.0,
    legs=None,
    chain_override=None,
):
    if legs is None:
        legs = [
            {
                "leg_index": 0, "token_in": "USDC", "token_out": "WETH",
                "protocol": "uniswap_v3", "venue": "uniswap_v3:base",
                "pool": "pool-a", "input_amount": 1000, "output_amount": 900,
                "quote": None, "fee_bps": 5, "quote_timestamp": None,
                "block": 10, "source": "uniswap_v3_quoter_base",
            },
            {
                "leg_index": 1, "token_in": "WETH", "token_out": "USDC",
                "protocol": "aerodrome_slipstream", "venue": "aerodrome:base",
                "pool": "pool-b", "input_amount": 900, "output_amount": 990,
                "quote": None, "fee_bps": 0, "quote_timestamp": None,
                "block": 11, "source": "aerodrome_quoter_base",
            },
        ]
    return {
        "ledger_id": f"ol1:SHADOW:{RUN}:{cid}",
        "opportunity_id": cid,
        "run_id": RUN,
        "candidate_id": cid,
        "verifier_bundle_id": f"flarb:{cid}" if bundle else None,
        "evidence": {
            "strategy": {
                "primary_family": family,
                "secondary_tags": ["CROSS_POOL"],
                "confidence": "MEDIUM",
                "evidence": ["stored evidence line"],
                "classifier_version": "phase0.strategy_intelligence.v1",
                "classification_completeness": "FULLY_CLASSIFIED",
                "classification_state": "COMPLETE",
            },
            "legs": legs,
            "economics": {
                "notional_usd": 10000.0 if bundle else None,
                "gross_spread_pct": -0.5 if bundle else None,
                "gross_profit_usd": gross_profit,
                "dex_fee_pct": 0.05 if bundle else None,
                "dex_fee_usd": None,
                "flash_loan_fee_usd": flash_fee if bundle else None,
                "flash_loan_fee_pct": None,
                "gas_cost_usd": 0.15 if bundle else None,
                "gas_units": 250000 if bundle else None,
                "gas_price": None,
                "slippage_pct": slippage if bundle else None,
                "slippage_usd": None,
                "mev_penalty": None,
                "mev_adjusted_net_usd": None,
                "mev_adjusted_net_pct": None,
                "total_cost_usd": None,
                "true_net_usd": net if bundle else None,
                "true_net_pct": None,
            },
            "gates": {
                "gate_7": {
                    "status": gate,
                    "reason": "atomic_profit $-59.31 < floor $25.00" if gate else None,
                },
                "gate_8": {"status": "NOT_EVALUATED", "reason": None},
                "gate_9": {"status": "NOT_EVALUATED", "reason": None},
                "final_observed_status": "denied:gate_rejection:gate_7:atomic_profit $-59.31 < floor $25.00",
                "decision_net_usd": net,
            },
            "provenance": {
                "chain": chain_override if chain_override is not None else chain,
                "flash_loan_provider": provider,
                "quote_source": "uniswap_v3_quoter_base" if bundle else None,
                "verifier_component": "flash_loan_arb_verifier" if bundle else None,
                "calculator_version": None,
                "classifier_version": "phase0.strategy_intelligence.v1",
                "git_sha": "823a79b617ddb1f19397cf5073b9c516aae4e9fd",
                "network_config_revision": "rev-1068cb9715194f118c99e5f04f9e1bdb",
                "rpc_identity": None,
                "timestamps": {
                    "hint_observed_at": verified_at - 10,
                    "verified_at": verified_at,
                    "bundle_created_at": "2026-10-05T05:30:00+00:00" if bundle else None,
                },
            },
        },
    }


def _population():
    rows = []
    plan = (
        ("DEX_TO_DEX", 17),
        ("CROSS_PROTOCOL", 8),
        ("CROSS_POOL", 27),
        ("MULTI_HOP", 27),
        ("COMPLEX_TRIANGULAR_CROSS_PROTOCOL", 83),
        ("MULTI_DEX", 34),
        ("TRIANGULAR", 92),
    )
    index = 0
    chains = ["ethereum", "arbitrum", "base", "optimism", "polygon", "bnb"]
    for family, count in plan:
        for _ in range(count):
            bundle = index < 180
            net = -59.31 if index == 0 else -100.0 - index
            rows.append(_row(
                f"c{index:03d}",
                family,
                net,
                chain=chains[index % 6],
                bundle=bundle,
                gate="FAIL" if bundle else None,
                verified_at=1791178200.0 + index,
            ))
            index += 1
    return rows


def test_summary_matches_certified_shape_without_source_collections():
    rows = _population()
    summary = summarize(rows)
    assert summary["total"] == 288
    assert summary["verifier_backed"] == 180
    assert summary["decision_only"] == 108
    assert summary["strategy_distribution"] == {
        "DEX_TO_DEX": 17,
        "CROSS_PROTOCOL": 8,
        "CROSS_POOL": 27,
        "MULTI_HOP": 27,
        "COMPLEX_TRIANGULAR_CROSS_PROTOCOL": 83,
        "MULTI_DEX": 34,
        "TRIANGULAR": 92,
    }
    assert summary["net"]["ge_0"] == 0
    assert summary["net"]["ge_25"] == 0
    assert summary["net"]["best"] == -59.31
    assert summary["best_opportunity"]["strategy_family"] == "DEX_TO_DEX"
    assert "hint_metric" not in summary


def test_list_filters_sort_and_pagination_omit_route_payload():
    rows = _population()
    base = apply_filters(rows, run_id=RUN, chain="base", bundle="backed", gate_7="FAIL")
    ordered = sort_rows(base, "net", "desc")
    page = paginate(ordered, 1, 25)
    assert page["total"] == len(base)
    assert len(page["items"]) <= 25
    assert page["items"][0]["decision_net_usd"] >= page["items"][-1]["decision_net_usd"]
    assert "legs" not in page["items"][0]
    assert "economics" not in page["items"][0]
    decision = apply_filters(rows, run_id=RUN, bundle="decision_only")
    assert len(decision) == 108
    assert all(row["verifier_bundle_id"] is None for row in decision)
    narrow = apply_filters(rows, run_id=RUN, family="DEX_TO_DEX", net_min=-60, net_max=-59)
    assert len(narrow) == 1
    assert narrow[0]["evidence"]["strategy"]["primary_family"] == "DEX_TO_DEX"


def test_detail_preserves_null_zero_gate_text_and_leg_order():
    row = _row("cand-dex", "DEX_TO_DEX", -59.31)
    full = detail(row)
    econ = full["evidence"]["economics"]
    assert econ["slippage_pct"] == 0.0
    assert econ["flash_loan_fee_usd"] == 0.0
    assert econ["gross_profit_usd"] is None
    assert econ["total_cost_usd"] is None
    assert econ["true_net_pct"] is None
    assert full["evidence"]["gates"]["gate_7"]["reason"] == "atomic_profit $-59.31 < floor $25.00"
    assert [leg["leg_index"] for leg in full["evidence"]["legs"]] == [0, 1]
    assert [leg["token_in"] for leg in full["evidence"]["legs"]] == ["USDC", "WETH"]
    slim = slim_row(row)
    assert slim["bundle_status"] == "verifier_bundle"
    assert "legs" not in slim


def test_timeline_omits_missing_timestamps():
    row = _row("cand-do", "MULTI_HOP", -10, bundle=False, gate=None)
    stages = [item["stage"] for item in timeline_rows(row)]
    assert stages == ["DISCOVERED", "VERIFIED"]
    assert "CLASSIFIED" not in stages
    assert "EXECUTED" not in stages


def test_sort_by_family_and_chain():
    rows = _population()[:6]
    families = [row["evidence"]["strategy"]["primary_family"] for row in sort_rows(rows, "family", "asc")]
    assert families == sorted(families)
    chains = [row["evidence"]["provenance"]["chain"] for row in sort_rows(rows, "chain", "asc")]
    assert chains == sorted(chains)


def test_workbook_sheets_counts_zero_null_and_secrets():
    rows = _population()
    poisoned = _row(
        "secret-row",
        "DEX_TO_DEX",
        -1,
        chain_override="https://eth-mainnet.g.alchemy.com/v2/supersecretvalue",
    )
    poisoned["evidence"]["provenance"]["quote_source"] = "mongodb://user:supersecretvalue@mongo/db"
    payload = build_workbook(rows + [poisoned])
    wb = load_workbook(io.BytesIO(payload))
    assert tuple(wb.sheetnames) == tuple(SHEETS)
    opportunities = wb["Opportunities"]
    # header + 288 + poisoned
    assert opportunities.max_row == 290
    economics = wb["Economics"]
    headers = [cell.value for cell in economics[1]]
    slip_col = headers.index("slippage_pct") + 1
    gross_col = headers.index("gross_profit_usd") + 1
    assert economics.cell(2, slip_col).value == 0
    assert economics.cell(2, gross_col).value is None
    path = wb["Price Path"]
    path_headers = [cell.value for cell in path[1]]
    idx_col = path_headers.index("leg_index") + 1
    token_col = path_headers.index("token_in") + 1
    first_indexes = [path.cell(row, idx_col).value for row in range(2, 4)]
    first_tokens = [path.cell(row, token_col).value for row in range(2, 4)]
    assert first_indexes == [0, 1]
    assert first_tokens == ["USDC", "WETH"]
    strategy = wb["Strategy Intelligence"]
    assert strategy.max_row == 290
    assert strategy.cell(2, 3).value == "DEX_TO_DEX"
    blob = " ".join(
        str(cell.value)
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "supersecretvalue" not in blob
    assert not SECRET.search(blob)
    assert wb["Summary"].cell(2, 2).value == RUN
    assert wb["Learning Dataset"].max_row == 290
    assert wb["Timeline"].max_row > 2
