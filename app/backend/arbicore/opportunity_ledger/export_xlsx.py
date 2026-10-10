"""Excel workbook built only from Opportunity Ledger rows."""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook

from .read_model import (
    bundle_status,
    chain_of,
    decision_net,
    detail,
    family_of,
    gate_reason,
    gate_status,
    provider_of,
    summarize,
    timeline_rows,
    _economics,
    _evidence,
    _gates,
    _provenance,
    _strategy,
)

SHEETS = (
    "Summary",
    "Opportunities",
    "Price Path",
    "Economics",
    "Strategy Intelligence",
    "Timeline",
    "Learning Dataset",
)


def _cell(value: Any) -> Any:
    """Blank means unavailable. Numeric zero stays zero."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value)
    return str(value)


def _write(ws, headers: List[str], rows: List[List[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_cell(value) for value in row])


def _legs(row: Dict[str, Any]) -> List[Dict[str, Any]]:
    legs = _evidence(row).get("legs")
    return [leg for leg in legs if isinstance(leg, dict)] if isinstance(legs, list) else []


def build_workbook(rows: List[Dict[str, Any]]) -> bytes:
    """Serialize ledger rows. Does not read any other collection."""
    cleaned = [detail(row) for row in rows]
    summary = summarize(cleaned)
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    _write(summary_ws, ["field", "value"], _summary_lines(summary))

    opp_headers = [
        "opportunity_id", "ledger_id", "run_id", "timestamp", "chain",
        "strategy_family", "secondary_tags", "provider", "bundle_status",
        "decision_net", "net_pct", "gate_7", "gate_7_reason", "gate_8",
        "gate_9", "final_status",
    ]
    opp_rows = []
    for row in cleaned:
        econ = _economics(row)
        opp_rows.append([
            row.get("opportunity_id"),
            row.get("ledger_id"),
            row.get("run_id"),
            row.get("display_timestamp"),
            chain_of(row),
            family_of(row),
            _strategy(row).get("secondary_tags"),
            provider_of(row),
            bundle_status(row),
            decision_net(row),
            econ.get("true_net_pct"),
            gate_status(row, "gate_7"),
            gate_reason(row, "gate_7"),
            gate_status(row, "gate_8"),
            gate_status(row, "gate_9"),
            _gates(row).get("final_observed_status"),
        ])
    opp_ws = wb.create_sheet("Opportunities")
    _write(opp_ws, opp_headers, opp_rows)

    path_headers = [
        "opportunity_id", "run_id", "leg_index", "token_in", "token_out",
        "protocol", "venue", "pool", "pool_address", "input_amount",
        "output_amount", "quote", "fee_bps", "quote_status",
        "quote_timestamp", "block", "source",
    ]
    path_rows = []
    for row in cleaned:
        for leg in _legs(row):
            path_rows.append([
                row.get("opportunity_id"),
                row.get("run_id"),
                leg.get("leg_index"),
                leg.get("token_in"),
                leg.get("token_out"),
                leg.get("protocol"),
                leg.get("venue"),
                leg.get("pool"),
                leg.get("pool_address"),
                leg.get("input_amount"),
                leg.get("output_amount"),
                leg.get("quote"),
                leg.get("fee_bps"),
                leg.get("quote_status"),
                leg.get("quote_timestamp"),
                leg.get("block"),
                leg.get("source"),
            ])
    path_ws = wb.create_sheet("Price Path")
    _write(path_ws, path_headers, path_rows)

    econ_headers = [
        "opportunity_id", "run_id", "notional_usd", "gross_spread_pct",
        "gross_profit_usd", "dex_fee_pct", "dex_fee_usd", "flash_loan_fee_usd",
        "flash_loan_fee_pct", "gas_cost_usd", "gas_units", "gas_price",
        "slippage_pct", "slippage_usd", "mev_penalty", "mev_adjusted_net_usd",
        "mev_adjusted_net_pct", "total_cost_usd", "true_net_usd", "true_net_pct",
        "decision_net_usd",
    ]
    econ_rows = []
    for row in cleaned:
        econ = _economics(row)
        econ_rows.append([
            row.get("opportunity_id"), row.get("run_id"),
            econ.get("notional_usd"), econ.get("gross_spread_pct"),
            econ.get("gross_profit_usd"), econ.get("dex_fee_pct"),
            econ.get("dex_fee_usd"), econ.get("flash_loan_fee_usd"),
            econ.get("flash_loan_fee_pct"), econ.get("gas_cost_usd"),
            econ.get("gas_units"), econ.get("gas_price"),
            econ.get("slippage_pct"), econ.get("slippage_usd"),
            econ.get("mev_penalty"), econ.get("mev_adjusted_net_usd"),
            econ.get("mev_adjusted_net_pct"), econ.get("total_cost_usd"),
            econ.get("true_net_usd"), econ.get("true_net_pct"),
            decision_net(row),
        ])
    econ_ws = wb.create_sheet("Economics")
    _write(econ_ws, econ_headers, econ_rows)

    strat_headers = [
        "opportunity_id", "run_id", "primary_family", "secondary_tags",
        "confidence", "evidence", "classification_completeness",
        "classification_state", "classifier_version",
    ]
    strat_rows = []
    for row in cleaned:
        strategy = _strategy(row)
        strat_rows.append([
            row.get("opportunity_id"),
            row.get("run_id"),
            strategy.get("primary_family"),
            strategy.get("secondary_tags"),
            strategy.get("confidence"),
            strategy.get("evidence"),
            strategy.get("classification_completeness"),
            strategy.get("classification_state"),
            strategy.get("classifier_version"),
        ])
    strat_ws = wb.create_sheet("Strategy Intelligence")
    _write(strat_ws, strat_headers, strat_rows)

    time_headers = ["opportunity_id", "ledger_id", "run_id", "stage", "timestamp"]
    time_rows = []
    for row in cleaned:
        for stage in timeline_rows(row):
            time_rows.append([
                stage["opportunity_id"], stage["ledger_id"], stage["run_id"],
                stage["stage"], stage["timestamp"],
            ])
    time_ws = wb.create_sheet("Timeline")
    _write(time_ws, time_headers, time_rows)

    learn_headers = [
        "opportunity_id", "run_id", "strategy_family", "chain", "provider",
        "leg_count", "stored_protocols", "stored_token_path",
        "decision_net_usd", "true_net_usd", "true_net_pct", "notional_usd",
        "gross_spread_pct", "gross_profit_usd", "dex_fee_usd",
        "flash_loan_fee_usd", "gas_cost_usd", "slippage_pct", "mev_penalty",
        "total_cost_usd", "gate_7", "gate_7_reason", "gate_8", "gate_8_reason",
        "gate_9", "gate_9_reason", "final_observed_status",
        "confidence", "classification_completeness", "classifier_version",
    ]
    learn_rows = []
    for row in cleaned:
        legs = _legs(row)
        econ = _economics(row)
        strategy = _strategy(row)
        learn_rows.append([
            row.get("opportunity_id"),
            row.get("run_id"),
            family_of(row),
            chain_of(row),
            provider_of(row),
            len(legs),
            [leg.get("protocol") for leg in legs],
            [leg.get("token_in") for leg in legs],
            decision_net(row),
            econ.get("true_net_usd"),
            econ.get("true_net_pct"),
            econ.get("notional_usd"),
            econ.get("gross_spread_pct"),
            econ.get("gross_profit_usd"),
            econ.get("dex_fee_usd"),
            econ.get("flash_loan_fee_usd"),
            econ.get("gas_cost_usd"),
            econ.get("slippage_pct"),
            econ.get("mev_penalty"),
            econ.get("total_cost_usd"),
            gate_status(row, "gate_7"),
            gate_reason(row, "gate_7"),
            gate_status(row, "gate_8"),
            gate_reason(row, "gate_8"),
            gate_status(row, "gate_9"),
            gate_reason(row, "gate_9"),
            _gates_final(row),
            strategy.get("confidence"),
            strategy.get("classification_completeness"),
            strategy.get("classifier_version"),
        ])
    learn_ws = wb.create_sheet("Learning Dataset")
    _write(learn_ws, learn_headers, learn_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _gates_final(row: Dict[str, Any]) -> Optional[str]:
    value = _gates(row).get("final_observed_status")
    return value if isinstance(value, str) else None


def _summary_lines(summary: Dict[str, Any]) -> List[List[Any]]:
    net = summary.get("net") or {}
    period = summary.get("period") or {}
    best = summary.get("best_opportunity") or {}
    worst = summary.get("worst_opportunity") or {}
    lines = [
        ["run_id", summary.get("run_id")],
        ["period_first_verified_at", period.get("first_verified_at_iso")],
        ["period_last_verified_at", period.get("last_verified_at_iso")],
        ["period_source", period.get("source")],
        ["total_opportunities", summary.get("total")],
        ["verifier_backed", summary.get("verifier_backed")],
        ["decision_only", summary.get("decision_only")],
        ["net_mean", net.get("mean")],
        ["net_best", net.get("best")],
        ["net_worst", net.get("worst")],
        ["net_ge_0", net.get("ge_0")],
        ["net_ge_25", net.get("ge_25")],
        ["best_opportunity_id", best.get("opportunity_id")],
        ["worst_opportunity_id", worst.get("opportunity_id")],
        ["blank_cell_means", "unavailable"],
        ["numeric_zero_means", "stored zero"],
    ]
    for label, key in (
        ("strategy", "strategy_distribution"),
        ("chain", "chain_distribution"),
        ("provider", "provider_distribution"),
        ("gate_7_status", "gate_7_status"),
    ):
        for name, count in sorted((summary.get(key) or {}).items()):
            lines.append([f"{label}:{name}", count])
    return lines
